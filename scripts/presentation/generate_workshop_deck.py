#!/usr/bin/env python3
"""Generate the editable Agentic AI Developer workshop PowerPoint deck."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "slides/github-agentic-ai-developer-workshop.pptx"
ASSETS = Path(__file__).with_name("assets")

WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)

BG = RGBColor(255, 255, 255)
PANEL = RGBColor(246, 248, 250)
PANEL_2 = RGBColor(246, 248, 250)
TEXT = RGBColor(36, 41, 46)
MUTED = RGBColor(88, 96, 105)
GREEN = RGBColor(34, 134, 58)
BLUE = RGBColor(3, 102, 214)
SECTION_BLUE = RGBColor(22, 88, 197)
PURPLE = RGBColor(111, 66, 193)
ORANGE = RGBColor(115, 92, 15)
RED = RGBColor(203, 36, 49)
BORDER = RGBColor(209, 213, 218)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)

TITLE_FONT = "Helvetica Neue"
BODY_FONT = "Helvetica Neue"
CODE_FONT = "Roboto Mono"


def apply_github_theme(prs):
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    theme = parse_xml(theme_part.blob)
    theme.set("name", "GitHub Developer Training")
    palette = {
        "dk1": "24292E", "lt1": "FFFFFF", "dk2": "586069", "lt2": "F6F8FA",
        "accent1": "0366D6", "accent2": "00A89D", "accent3": "22863A",
        "accent4": "735C0F", "accent5": "CB2431", "accent6": "6F42C1",
        "hlink": "0366D6", "folHlink": "6F42C1",
    }
    for name, value in palette.items():
        color = theme.xpath(f"./a:themeElements/a:clrScheme/a:{name}")[0]
        for child in list(color):
            color.remove(child)
        rgb = OxmlElement("a:srgbClr")
        rgb.set("val", value)
        color.append(rgb)
    for font in theme.xpath("./a:themeElements/a:fontScheme/*/a:latin"):
        font.set("typeface", BODY_FONT)
    for effects in theme.xpath("./a:themeElements/a:fmtScheme/a:effectStyleLst/a:effectStyle/a:effectLst"):
        for effect in list(effects):
            effects.remove(effect)
    theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8",
                                      standalone=True)


def set_background(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=24, color=TEXT, bold=False,
             font=BODY_FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=21, color=TEXT, accent=BLUE,
                spacing=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.12)
    frame.margin_right = Inches(0.06)
    for index, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = level
        paragraph.font.name = BODY_FONT
        paragraph.font.size = Pt(size - level * 2)
        paragraph.font.color.rgb = color if level == 0 else MUTED
        paragraph.space_after = Pt(spacing)
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(Inches(0.22 + level * 0.18)))
        properties.set("indent", str(-Inches(0.18)))
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•" if level == 0 else "○")
        properties.append(bullet)
    return box


def add_rect(slide, x, y, w, h, fill=PANEL, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_card(slide, title, body, x, y, w, h, accent=BLUE, number=None):
    add_rect(slide, x, y, w, h)
    if number is not None:
        add_text(slide, str(number), x + 0.2, y + 0.2, 0.35, 0.55, 17, accent, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        title_x = x + 0.7
        title_w = w - 0.9
    else:
        title_x = x + 0.3
        title_w = w - 0.55
    title_height = 0.65
    title_size = 14 if number is not None and w < 3.2 else 18
    add_text(slide, title, title_x, y + 0.15, title_w, title_height, title_size, TEXT, True)
    body_y = y + 0.2 + title_height
    add_text(slide, body, x + 0.3, body_y, w - 0.55,
             h - title_height - 0.3, 16 if h >= 2.5 else 14, TEXT)


def add_code(slide, code, x, y, w, h, size=13, label=None):
    add_rect(slide, x, y, w, h)
    if label:
        add_text(slide, label.upper(), x + 0.2, y + 0.12, w - 0.4, 0.25, 10, BLUE, True)
        code_y = y + 0.45
        code_h = h - 0.58
    else:
        code_y = y + 0.18
        code_h = h - 0.3
    add_text(slide, code, x + 0.2, code_y, w - 0.4, code_h, size, TEXT, False,
             font=CODE_FONT, margin=0)


def add_header(slide, title, section=None, number=None):
    add_text(slide, title, 0.62, 0.42, 12.0, 1.15, 32, TEXT, True, TITLE_FONT,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_footer(slide, section, number, color=MUTED):
    add_text(slide, str(number), 12.55, 7.12, 0.4, 0.24, 9, color,
             align=PP_ALIGN.RIGHT)


def add_notes(slide, notes):
    slide.notes_slide.notes_text_frame.text = notes.strip()


def add_title_artwork(slide):
    set_background(slide, BLACK)
    slide.shapes.add_picture(str(ASSETS / "github-title-background.png"),
                             0, 0, width=WIDTH, height=HEIGHT)
    slide.shapes.add_picture(str(ASSETS / "github-mark-white.png"),
                             Inches(0.85), Inches(1.15), width=Inches(0.33))
    add_text(slide, "GitHub", 1.28, 1.14, 2.2, 0.33, 12, WHITE, font=CODE_FONT)


def base_slide(prs, title, section, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    number = len(prs.slides)
    add_header(slide, title, section, number)
    add_footer(slide, section, number)
    add_notes(slide, notes)
    return slide


def comparison_slide(prs, title, section, notes, first, second):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_rect(slide, 0, 0, 6.3, 7.5)
    add_text(slide, title, 0.65, 2.6, 5.0, 2.0, 34, TEXT, True, TITLE_FONT,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for y, (heading, items) in zip((1.0, 3.9), (first, second)):
        add_text(slide, heading, 7.0, y, 5.4, 0.55, 22, TEXT, True)
        add_bullets(slide, items, 7.0, y + 0.7, 5.5, 2.0, 18, spacing=6)
    add_footer(slide, section, len(prs.slides))
    add_notes(slide, notes)
    return slide


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_artwork(slide)
    add_text(slide, "Agentic AI\nDeveloper\nWorkshop", 0.85, 2.0, 10.8, 2.7, 52, WHITE, True,
             TITLE_FONT)
    add_text(slide, "From prompts to a governed, observable agentic SDLC", 0.9, 5.25,
             10.8, 0.65, 22, WHITE)
    add_text(slide, "6 modules  |  2 workflows  |  7 labs  |  1 capstone",
             0.9, 6.2, 11.0, 0.35, 14, WHITE, font=CODE_FONT)
    add_text(slide, "Workshop companion for GH-600", 0.9, 6.72, 8.0, 0.3, 12, WHITE,
             font=CODE_FONT)
    add_notes(slide, """Welcome everyone
We will be building and inspecting agent workflows today
Not trying to give an agent unlimited autonomy
We want to know what it can do, when it should stop, and how we check its work

Good time for a round of introductions
What is your experience with GitHub and Copilot?
What are you excited to learn today?""")


def section_slide(prs, module_number, title, subtitle, color, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, SECTION_BLUE)
    number = len(prs.slides)
    add_text(slide, f"Module {module_number}", 0.6, 0.72, 4.0, 0.4, 16, WHITE)
    add_text(slide, title, 0.6, 2.0, 11.85, 2.1, 46, WHITE, True, TITLE_FONT)
    add_text(slide, subtitle, 0.65, 5.1, 11.8, 1.3, 23, WHITE)
    add_footer(slide, f"Module {module_number}", number, WHITE)
    add_notes(slide, notes)


def flow(slide, labels, x, y, total_w, node_h=0.95, colors=None, arrows=True):
    gap = 0.32
    node_w = (total_w - gap * (len(labels) - 1)) / len(labels)
    colors = colors or [BLUE] * len(labels)
    centers = []
    for index, label in enumerate(labels):
        node_x = x + index * (node_w + gap)
        add_rect(slide, node_x, y, node_w, node_h, fill=PANEL_2, line=colors[index])
        label_size = 10 if node_w < 1.1 else 12 if node_w < 1.8 else 16
        add_text(slide, label, node_x + 0.05, y + 0.12, node_w - 0.1, node_h - 0.2,
                 label_size, TEXT, True, align=PP_ALIGN.CENTER,
                 valign=MSO_ANCHOR.MIDDLE, margin=0)
        centers.append((node_x + node_w / 2, y + node_h / 2))
        if arrows and index:
            previous_right = x + (index - 1) * (node_w + gap) + node_w
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(previous_right + 0.04), Inches(y + node_h / 2),
                Inches(node_x - 0.04), Inches(y + node_h / 2),
            )
            connector.line.color.rgb = MUTED
            connector.line.width = Pt(1.5)
    return centers


def make_deck(output=OUTPUT):
    prs = Presentation()
    apply_github_theme(prs)
    prs.slide_width = WIDTH
    prs.slide_height = HEIGHT
    prs.core_properties.title = "GitHub Certified: Agentic AI Developer Workshop"
    prs.core_properties.subject = "Instructor-ready workshop deck aligned to GH-600"
    prs.core_properties.author = "Agentic SDLC Workshop"
    prs.core_properties.keywords = "GitHub, Copilot, agentic AI, GH-600, SDLC, MCP"

    title_slide(prs)

    slide = base_slide(prs, "What learners will be able to do", "Workshop outcomes",
                       """Hands-on training
Each module gives us something we can inspect
A contract, a plan, a tool decision, or validation evidence

Ask which area is most relevant to their day-to-day work
Keep that example in mind as we go through the labs""")
    outcomes = [
        ("BOUND", "Define goals, tools, state, policies, stop conditions, and escalation."),
        ("CONNECT", "Exchange versioned artifacts between planning and implementation."),
        ("CONTROL", "Apply allowlists, least privilege, approvals, and protected paths."),
        ("ORCHESTRATE", "Coordinate specialists with concurrency, fan-in, and recovery."),
        ("EVALUATE", "Persist state and enforce deterministic quality gates."),
        ("OPERATE", "Audit, observe, roll back, and improve agent workflows."),
    ]
    for i, (head, body) in enumerate(outcomes):
        col, row = i % 3, i // 3
        add_card(slide, head, body, 0.68 + col * 4.16, 1.78 + row * 2.35, 3.8, 1.95,
                 [BLUE, PURPLE, GREEN, ORANGE, BLUE, RED][i])

    slide = base_slide(prs, "A 4 hour 40 minute learning journey", "Agenda",
                       """We will work through six modules, then a capstone
Share the lab links before starting

Module 4 is the longer block
We run the parallel workflow and spend 45 minutes on Project Pulse
Make sure everyone has time to run the app, not just generate files""")
    agenda = [
        ("00:00", "Foundations", "20 min"), ("00:20", "Architecture", "35 min"),
        ("00:55", "Tools + MCP", "40 min"), ("01:35", "Orchestration", "75 min"),
        ("02:50", "Memory + evaluation", "40 min"), ("03:30", "Governance", "40 min"),
        ("04:10", "Capstone", "30 min"),
    ]
    y = 1.72
    for index, (time, label, duration) in enumerate(agenda):
        color = [BLUE, PURPLE, ORANGE, GREEN, BLUE, RED, PURPLE][index]
        add_text(slide, time, 0.75, y, 1.0, 0.38, 15, color, True)
        add_rect(slide, 1.75, y - 0.02, 9.7, 0.53, fill=PANEL, line=BORDER)
        add_text(slide, label, 2.0, y + 0.07, 6.4, 0.28, 16, TEXT, True)
        add_text(slide, duration, 9.55, y + 0.07, 1.5, 0.28, 14, MUTED, True, align=PP_ALIGN.RIGHT)
        y += 0.72

    slide = base_slide(prs, "GH-600 certification domains", "Certification map",
                       """These are weighting ranges, not exact question counts
We will connect the domains to things we can run in the repository

Architecture first, then tools, coordination, memory, and governance
The important part is understanding why each control is there""")
    domains = [
        ("Architecture + SDLC", "15–20%", BLUE), ("Tools + environment", "20–25%", PURPLE),
        ("Memory + state", "10–15%", ORANGE), ("Evaluation + tuning", "15–20%", GREEN),
        ("Multi-agent coordination", "15–20%", BLUE), ("Guardrails + accountability", "10–15%", RED),
    ]
    for i, (name, weight, color) in enumerate(domains):
        col, row = i % 2, i // 2
        x, y = 0.78 + col * 6.2, 1.75 + row * 1.67
        add_rect(slide, x, y, 5.75, 1.28, fill=PANEL)
        add_text(slide, name, x + 0.25, y + 0.22, 3.9, 0.35, 18, TEXT, True)
        add_text(slide, weight, x + 4.35, y + 0.19, 1.05, 0.4, 18, color, True, align=PP_ALIGN.RIGHT)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.25), Inches(y + 0.83), Inches(5.0), Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = BORDER; bar.line.fill.background()
        value = float(weight.split("–")[1].replace("%", "")) / 25
        active = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.25), Inches(y + 0.83), Inches(5.0 * value), Inches(0.12))
        active.fill.solid(); active.fill.fore_color.rgb = color; active.line.fill.background()

    slide = base_slide(prs, "The six-module journey", "Learning path",
                       """Start with the vocabulary
Then define a contract and decide which tools the agent can use
After that we can coordinate agents and keep useful state

Governance is not something we add at the end
The earlier decisions give us the evidence we need to review and recover""")
    flow(slide, ["1\nFoundations", "2\nArchitecture", "3\nTools", "4\nOrchestration", "5\nMemory", "6\nGovernance"],
         0.72, 2.18, 11.9, 1.25, [BLUE, PURPLE, ORANGE, GREEN, BLUE, RED])
    add_text(slide, "Vocabulary", 0.9, 4.25, 1.4, 0.35, 14, MUTED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Contracts", 2.95, 4.25, 1.4, 0.35, 14, MUTED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Authority", 4.95, 4.25, 1.4, 0.35, 14, MUTED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Coordination", 6.95, 4.25, 1.6, 0.35, 14, MUTED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Continuity", 9.0, 4.25, 1.4, 0.35, 14, MUTED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Accountability", 10.9, 4.25, 1.7, 0.35, 14, MUTED, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 1.2, 5.35, 10.9, 0.78, fill=PANEL_2, line=GREEN)
    add_text(slide, "Every module produces an artifact that the next module can inspect.",
             1.45, 5.57, 10.4, 0.35, 20, TEXT, True, align=PP_ALIGN.CENTER)

    section_slide(prs, 1, "Foundations of Agentic AI", "Move from one-shot assistance to bounded systems that plan, act, observe, evaluate, and escalate.", BLUE,
                  """First, what do we mean by an agent?
Not a personality or a chat window
Look at the loop, the tools, and the authority it has""")

    comparison_slide(prs, "Assistant versus agent", "Module 1 — Foundations",
                     """An assistant can suggest a change
An agent may be able to make that change and continue working
That is where the risk changes

Ask what this system can do without another human decision
Tool calls and changes to state are the things to look for""",
                     ("Assistant", ["Responds to a prompt", "Proposes content",
                                    "Human initiates each step", "Usually has limited state"]),
                     ("Agent", ["Pursues a goal and calls tools", "Observes results",
                                "Maintains state", "Stops or escalates by policy"]))

    slide = base_slide(prs, "The bounded agent lifecycle", "Module 1 — Foundations",
                       """Walk through the loop from left to right
We give the agent a goal, it plans, acts, and checks what happened
Evaluation decides what happens next

Stop, retry, re-plan, or ask a human
A message saying 'done' is not enough
Show the artifact and the check that passed""")
    flow(slide, ["GOAL", "PLAN", "ACT", "OBSERVE", "EVALUATE"], 0.8, 2.1, 11.7, 1.08,
         [BLUE, PURPLE, ORANGE, BLUE, GREEN])
    add_rect(slide, 3.2, 4.15, 6.9, 1.05, fill=PANEL_2, line=RED)
    add_text(slide, "STOP / RETRY / RE-PLAN / ESCALATE", 3.45, 4.46, 6.4, 0.35, 19, RED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Policy surrounds the loop: scope • permissions • timeout • evidence • human control",
             1.2, 5.75, 10.9, 0.45, 17, MUTED, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "GitHub is the system of record and control plane", "Module 1 — Foundations",
                       """GitHub is where we keep the work and the evidence
An issue tells us what was requested
A pull request shows what changed
Checks tell us what passed, and rulesets control what can merge

Ask where learners currently record an approval
Could another person find it later?""")
    cards = [
        ("ISSUE", "Intent + acceptance criteria", BLUE), ("BRANCH", "Isolated execution", PURPLE),
        ("PULL REQUEST", "Reviewable change set", GREEN), ("CHECK", "Deterministic evidence", ORANGE),
        ("ARTIFACT", "Durable handoff", BLUE), ("RULESET", "Enforced policy", RED),
    ]
    for i, (head, body, color) in enumerate(cards):
        col, row = i % 3, i // 3
        add_card(slide, head, body, 0.72 + col * 4.18, 1.78 + row * 2.25, 3.84, 1.8, color)

    slide = base_slide(prs, "Agent contract: make authority explicit", "Module 1 — Foundations",
                       """Before looking at the JSON, explain the fields
What is the goal?
Which tools can it use, and what is it not allowed to do?
When does it stop or ask for help?

Open Lab 1 and write the contract
Check everyone has a stop condition and an escalation rule
Those are easy to leave out""")
    contract = '''{
  "objective": "Generate a project health report",
  "inputs": ["project metadata"],
  "allowed_tools": ["repository.read", "artifact.write"],
  "forbidden": ["push to protected branch", "read secrets"],
  "success": ["valid JSON", "SHA-256 recorded"],
  "timeout_minutes": 10,
  "escalate_when": ["risk is high", "validation fails"]
}'''
    add_code(slide, contract, 0.75, 1.75, 7.15, 4.95, 13, "agent-contract.json")
    add_card(slide, "LAB 1", "Write a bounded contract, generate a plan artifact, and identify one field that would be unsafe to leave implicit.", 8.35, 1.75, 4.2, 2.1, BLUE)
    add_card(slide, "COMPLETION EVIDENCE", "Point to objective, tools, stop condition, and escalation in an inspectable artifact.", 8.35, 4.18, 4.2, 2.05, GREEN)

    section_slide(prs, 2, "Agent Architecture and SDLC Integration", "Separate responsibilities and connect them through versioned, validated handoff contracts.", PURPLE,
                  """Now we have the vocabulary, we can connect the pieces
One process needs to pick up another process's work
It should not have to guess what the output means""")

    slide = base_slide(prs, "Single responsibility, explicit ownership", "Module 2 — Architecture",
                       """Each role has a job and an output it owns
The Planner writes the plan
The Implementer makes the payload
The Validator checks it
The human owns the approval decision

If something fails, we know where to look
File ownership also helps keep two agents from editing the same thing""")
    roles = [
        ("PLANNER", "Turns intent into ordered work\nOwns plan.json", PURPLE),
        ("IMPLEMENTER", "Creates requested payload\nOwns health-report.json", BLUE),
        ("VALIDATOR", "Checks contract + evidence\nOwns validation.json", GREEN),
        ("HUMAN REVIEWER", "Approves consequential changes\nOwns the decision", ORANGE),
    ]
    for i, (head, body, color) in enumerate(roles):
        col, row = i % 2, i // 2
        add_card(slide, head, body, 0.8 + col * 6.15, 1.8 + row * 2.35, 5.55, 1.95, color, i + 1)

    slide = base_slide(prs, "A versioned plan is an API", "Module 2 — Architecture",
                       """Treat the plan like an API between jobs
The schema version tells the next job what it can read
The request hash tells us which input produced it
Acceptance criteria go with the work

Show what happens if a required field is missing
The implementer should reject that plan, not guess a default""")
    plan = '''{
  "schema_version": "1.0",
  "request_id": "demo-plan-001",
  "request_hash": "16c8…d421",
  "project_metadata": {
    "name": "agentic-sdlc-demo",
    "status": "healthy",
    "owner": "platform-engineering"
  },
  "acceptance_criteria": [
    "The report contains project name and status",
    "The report is valid JSON"
  ],
  "handoff": {
    "producer": "deterministic-planner",
    "consumer": "deterministic-implementer"
  }
}'''
    add_code(slide, plan, 0.72, 1.7, 7.7, 5.15, 11.5, "plan.json")
    add_bullets(slide, ["Stable schema version", "Canonical request hash", "Explicit producer + consumer", "Acceptance criteria travel with work", "Risk determines human approval"], 8.78, 1.9, 3.8, 4.6, 17)

    slide = base_slide(prs, "Concrete plan → implement handoff", "Module 2 — Architecture",
                       """Open .github/workflows/plan-implement.yml
Follow the plan artifact between the jobs
The first job uploads plan.json
The next job downloads it and runs apply_plan.py

Run the lab and open the generated health report
There should also be a manifest with the digest and byte size
Make sure everyone can find both outputs""")
    workflow = '''jobs:
  plan:
    runs-on: ubuntu-latest
    steps:
      - uses: actions/checkout@v4
      - run: |
          python3 scripts/plan/generate_plan.py \\
            --request examples/plan_request.json \\
            --output artifacts/plan/plan.json
      - uses: actions/upload-artifact@v4
        with:
          name: agent-plan
          path: artifacts/plan/plan.json

  implement:
    needs: plan
    runs-on: ubuntu-latest
    steps:
      - uses: actions/download-artifact@v4
        with:
          name: agent-plan
          path: artifacts/plan
      - run: |
          python3 scripts/implement/apply_plan.py \\
            --plan artifacts/plan/plan.json \\
            --output artifacts/implementation/implementation.json'''
    add_code(slide, workflow, 0.72, 1.66, 8.35, 5.3, 10.4, ".github/workflows/plan-implement.yml")
    add_card(slide, "NOT A PLACEHOLDER", "The consumer writes health-report.json and records its SHA-256 digest and byte size.", 9.42, 1.75, 3.2, 2.2, GREEN)
    add_card(slide, "BOUNDARY", "Jobs exchange immutable artifacts instead of hidden process memory.", 9.42, 4.22, 3.2, 2.0, PURPLE)

    slide = base_slide(prs, "Evidence: payload + integrity manifest", "Module 2 — Architecture",
                       """The report is the payload
The manifest records which bytes we handed off

Run sha256sum on the generated report
Compare it with the manifest
Now change the report and run it again
The digest no longer matches
That is how the next job can detect a changed artifact""")
    flow(slide, ["request.json", "plan.json", "health-report.json", "implementation.json"],
         0.75, 2.0, 11.85, 1.15, [BLUE, PURPLE, GREEN, ORANGE])
    add_code(slide, '''"generated_files": [
  {
    "path": "health-report.json",
    "sha256": "cc7e6b…ec7c",
    "bytes": 742
  }
]''', 1.0, 4.0, 5.4, 2.2, 14, "implementation.json")
    add_bullets(slide, ["Artifact exists", "JSON parses", "Required fields are present", "Digest matches exact bytes", "Protected paths remain untouched"], 7.0, 3.95, 5.0, 2.4, 18)

    slide = base_slide(prs, "Pull-request governance belongs in the architecture", "Module 2 — Architecture",
                       """Before giving an agent write access, decide how its work reaches main
What checks must pass?
Who needs to review it?
Can it bypass a protected branch?

Use the same pull request controls we would use for a teammate
Have a recovery path before a deployment, not after it fails""")
    controls = [
        ("TEMPLATE", "Required context and test plan"), ("CHECKS", "Objective machine evidence"),
        ("CODEOWNERS", "Domain-specific human review"), ("RULESET", "No bypass of protected branch"),
        ("ENVIRONMENT", "Approval before deployment"), ("ROLLBACK", "Known recovery reference"),
    ]
    for i, (head, body) in enumerate(controls):
        col, row = i % 3, i // 3
        add_card(slide, head, body, 0.72 + col * 4.18, 1.8 + row * 2.25, 3.84, 1.78,
                 [BLUE, GREEN, PURPLE, RED, ORANGE, BLUE][i])

    section_slide(prs, 3, "Tooling, MCP, and Execution Environments", "Expose narrow capabilities through explicit protocol, authorization, and runtime boundaries.", ORANGE,
                  """Next we will connect tools
Being able to call a tool does not mean every action is allowed
MCP handles the connection
Our policy still decides what the agent can do""")

    slide = base_slide(prs, "A tool boundary has four layers", "Module 3 — Tools + MCP",
                       """Start with the request shape
Then the named tool and operation
Then check the permissions and path
Finally record the decision and the reason

A valid JSON-RPC request can still be denied
Correct format is not the same thing as permission""")
    layers = [
        ("1", "PROTOCOL", "Shape, identity, arguments", BLUE),
        ("2", "CAPABILITY", "Named tool + operation", PURPLE),
        ("3", "AUTHORIZATION", "Allowlists + path scope", ORANGE),
        ("4", "EVIDENCE", "Decision + reason + provenance", GREEN),
    ]
    for i, (num, head, body, color) in enumerate(layers):
        add_card(slide, head, body, 1.0 + i * 3.02, 2.0, 2.65, 2.55, color, num)
    add_text(slide, "Protocol compatibility ≠ permission to act", 1.2, 5.45, 10.9, 0.55, 24, RED, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "Concrete MCP-style request and decision", "Module 3 — Tools + MCP",
                       """Run the read request first
Look at the allowed decision

Change the operation to write_file and run it again
Same request format, different permission outcome
Check the nonzero exit code and the denied audit record
Field any questions about the difference between a protocol and a policy""")
    request = '''{
  "jsonrpc": "2.0",
  "id": "mcp-demo-001",
  "tool": "repository.read",
  "operation": "read_file",
  "path": "examples/project_spec.json",
  "arguments": {"encoding": "utf-8"}
}'''
    denied = '''{
  "decision": "deny",
  "reason": "operation is not allowed",
  "tool": "repository.read",
  "operation": "write_file",
  "path": "examples/project_spec.json"
}'''
    add_code(slide, request, 0.72, 1.72, 5.85, 4.8, 13, "request.json")
    add_code(slide, denied, 6.88, 1.72, 5.72, 4.8, 13, "denied response")

    comparison_slide(prs, "Choose the execution environment deliberately",
                     "Module 3 — Tools + MCP",
                     """We use Codespaces when a person is exploring and inspecting the work
Actions gives us a repeatable run with logs and artifacts
An agent workflow may use both

Check permissions in each environment
Network access, secrets, write scope, timeout, and retention
Do not assume the interactive session and the workflow have the same authority""",
                     ("Codespaces", ["Interactive, with a human present",
                                     "Rich repository context", "Good for exploration",
                                     "Scoped CLI permissions"]),
                     ("GitHub Actions", ["Repeatable and event-driven",
                                        "Explicit token permissions", "Artifacts and logs",
                                        "Environment approval gates"]))

    section_slide(prs, 4, "Multi-Agent Systems and Orchestration", "Coordinate specialist agents through isolated work, observable artifacts, fan-in, and safe recovery.", GREEN,
                  """We have two labs in this section
First the deterministic parallel workflow
Then the Project Pulse custom-agent team

More agents does not automatically mean better coordination
We need ownership, handoffs, and a way to handle failures""")

    slide = base_slide(prs, "Deterministic multi-agent topology", "Module 4 — Orchestration",
                       """Follow the arrows
The Spec Analyzer and Risk Reviewer can run at the same time
They have separate outputs

The merger waits for both reports
That is the fan-in point
Then a separate validator checks the merged plan
Ask what should happen if only one report arrives""")
    add_rect(slide, 0.85, 2.45, 2.25, 1.0, fill=PANEL_2, line=BLUE)
    add_text(slide, "ORCHESTRATOR", 1.0, 2.75, 1.95, 0.4, 14, TEXT, True,
             align=PP_ALIGN.CENTER, margin=0)
    add_rect(slide, 4.0, 1.6, 2.55, 1.0, fill=PANEL_2, line=PURPLE)
    add_text(slide, "SPEC ANALYZER", 4.2, 1.9, 2.15, 0.3, 16, TEXT, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 4.0, 3.55, 2.55, 1.0, fill=PANEL_2, line=ORANGE)
    add_text(slide, "RISK REVIEWER", 4.2, 3.85, 2.15, 0.3, 16, TEXT, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 7.45, 2.45, 2.2, 1.0, fill=PANEL_2, line=GREEN)
    add_text(slide, "PLAN MERGER", 7.65, 2.75, 1.8, 0.3, 16, TEXT, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 10.55, 2.45, 1.9, 1.0, fill=PANEL_2, line=RED)
    add_text(slide, "VALIDATOR", 10.72, 2.75, 1.55, 0.3, 16, TEXT, True, align=PP_ALIGN.CENTER)
    for x1, y1, x2, y2 in [(3.1,2.95,4.0,2.1),(3.1,2.95,4.0,4.05),(6.55,2.1,7.45,2.95),(6.55,4.05,7.45,2.95),(9.65,2.95,10.55,2.95)]:
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = MUTED; line.line.width = Pt(2)
    add_text(slide, "parallel", 4.68, 5.05, 1.2, 0.3, 13, PURPLE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "fan-in", 7.85, 4.05, 1.3, 0.3, 13, GREEN, True, align=PP_ALIGN.CENTER)
    add_text(slide, "quality gate", 10.65, 4.05, 1.65, 0.3, 13, RED, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "Concrete parallel specialists and fan-in", "Module 4 — Orchestration",
                       """Open .github/workflows/multi-agent.yml
Show the two matrix roles and their Python executors
Each specialist gets its own artifact directory

The slide shows an excerpt; use the workflow file for the full commands
Run it and find both uploaded reports
Then inspect the merged plan and validation result
The concurrency group keeps merge jobs for the same branch from overlapping""")
    multi = '''jobs:
  specialists:
    strategy:
      fail-fast: false
      matrix:
        role: [spec-analyzer, risk-reviewer]
    steps:
      - uses: actions/checkout@v4
      - name: Run specialist
        run: |
          if [ "${{ matrix.role }}" = "spec-analyzer" ]; then
            python3 scripts/multi_agent/spec_analyzer.py \\
              --spec examples/project_spec.json \\
              --output artifacts/${{ matrix.role }}/report.json
          else
            python3 scripts/multi_agent/risk_reviewer.py \\
              --spec examples/project_spec.json \\
              --output artifacts/${{ matrix.role }}/report.json
          fi

  fan-in:
    needs: specialists
    concurrency:
      group: multiagent-${{ github.ref }}
    steps:
      - run: python3 scripts/multi_agent/plan_merger.py ...'''
    add_code(slide, multi, 0.72, 1.66, 8.7, 5.33, 10.5, ".github/workflows/multi-agent.yml")
    add_card(slide, "ISOLATION", "Separate artifact directory per specialist.", 9.73, 1.75, 2.85, 1.45, PURPLE)
    add_card(slide, "RECOVERY", "fail-fast: false preserves independent evidence.", 9.73, 3.45, 2.85, 1.45, ORANGE)
    add_card(slide, "SERIALIZATION", "One branch-scoped merge job at a time.", 9.73, 5.15, 2.85, 1.45, GREEN)

    slide = base_slide(prs, "Failures are orchestration outcomes", "Module 4 — Orchestration",
                       """Failures are part of the workflow
Keep the successful specialist's output if the other one fails
Do not discard useful evidence

Ask which of these can retry automatically
Which one needs a human?
Retries need a limit, and rerunning a step should not duplicate the side effects""")
    failures = [
        ("SPECIALIST FAILS", "Preserve other output\nRetry bounded times\nEscalate with logs", RED),
        ("OUTPUT IS STALE", "Record input hash\nReject mismatched provenance", ORANGE),
        ("AGENTS DISAGREE", "Use deterministic arbitration\nRoute unresolved conflict to human", PURPLE),
        ("MERGE OVERLAPS", "Concurrency group\nBranch isolation\nIdempotent fan-in", GREEN),
    ]
    for i, (head, body, color) in enumerate(failures):
        col, row = i % 2, i // 2
        add_card(slide, head, body, 0.8 + col * 6.15, 1.8 + row * 2.35, 5.55, 1.95, color)

    slide = base_slide(prs, "Project Pulse: build an AI dream team", "Module 4 — Required lab",
                       """Now we will build Mona's Project Pulse dashboard
This is the in-repository adaptation of the Microsoft Learn and GitHub Skills exercise
We will use Copilot CLI and the four custom agents

Open labs/04-multi-agent/project-pulse/README.md
Walk through the required files before starting
We have 45 minutes
Leave time to run the dashboard and inspect it""")
    add_text(slide, "45 MINUTES", 0.8, 1.78, 2.2, 0.4, 15, GREEN, True)
    add_text(slide, "Plan, design, build, run, and validate Mona's Project Pulse dashboard.", 0.8, 2.28, 11.2, 0.75, 26, TEXT, True)
    deliverables = ["docs/agent-team.md", "docs/project-pulse-plan.md", "app/index.html", "app/styles.css", "app/project-data.json", ".vscode/launch.json", "docs/final-handoff.md"]
    for i, item in enumerate(deliverables):
        col, row = i % 2, i // 2
        add_rect(slide, 0.85 + col * 6.05, 3.4 + row * 0.68, 5.55, 0.48, fill=PANEL_2)
        add_text(slide, item, 1.08 + col * 6.05, 3.51 + row * 0.68, 5.1, 0.25, 14, BLUE if i < 2 else TEXT, True, font=CODE_FONT)

    slide = base_slide(prs, "The custom agent team", "Module 4 — Project Pulse",
                       """Open .github/agents and read the definitions
The Orchestrator coordinates; it does not implement the app
The Planner sequences the work
The Designer owns the visual decisions
The Coder implements and checks runnable behavior

Look at tools, file scope, and handoff rules
Those boundaries matter more than just choosing a model""")
    team = [
        ("ORCHESTRATOR", "Coordinates phases\nDelegates; does not implement", GREEN),
        ("PLANNER", "Researches + sequences\nDefines ownership + dependencies", PURPLE),
        ("DESIGNER", "UX + accessibility\nOwns visual decisions", ORANGE),
        ("CODER", "Implements files\nValidates runnable behavior", BLUE),
    ]
    for i, (head, body, color) in enumerate(team):
        add_card(slide, head, body, 0.65 + i * 3.15, 1.95, 2.85, 3.15, color, i + 1)
    add_text(slide, "Orchestrator grants explicit file scope; specialists report files touched and validation evidence.", 1.0, 5.75, 11.2, 0.55, 18, MUTED, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "Project Pulse orchestration sequence", "Module 4 — Project Pulse",
                       """Select Orchestrator with /agent
Give it the outcome and the exact deliverables from the lab
The Planner goes first
Then design and implementation follow the agreed file boundaries

Run the app and validate it
If a check fails, send it back to the specialist who owns that file
Open the dashboard at a narrow and a wide viewport
Check the browser console before writing the handoff""")
    flow(slide, ["INSPECT TEAM", "PLAN", "DESIGN", "IMPLEMENT", "RUN", "VALIDATE", "HANDOFF"],
         0.55, 2.0, 12.25, 1.08, [BLUE, PURPLE, ORANGE, BLUE, GREEN, RED, PURPLE])
    add_card(slide, "PROMPT CONTRACT", "Name the desired outcome, exact files, required fields, launch behavior, and validation command.", 0.82, 4.0, 5.65, 1.75, BLUE)
    add_card(slide, "HUMAN EVIDENCE", "Inspect the rendered dashboard at narrow and wide viewports and verify a clean browser console.", 6.86, 4.0, 5.65, 1.75, GREEN)

    slide = base_slide(prs, "Project Pulse completion evidence", "Module 4 — Project Pulse",
                       """Run the validation command
Check everyone can explain what it actually checks
Files, JSON fields, source wiring, and launch configuration

It does not execute the browser JavaScript
A passing result is not proof that the dashboard works
Open the app, check the project cards, resize the window, and inspect the console
Record those checks in the final handoff""")
    add_code(slide, '''python3 scripts/project_pulse/validate_exercise.py --root .

PASS: project data has at least two projects
PASS: dashboard source fetches project-data.json
PASS: launch configuration is runnable
PASS: launch configuration opens index.html
PASS: final handoff names every agent and deliverable
All Project Pulse validation checks passed.''', 0.72, 1.72, 7.4, 4.85, 13, "terminal")
    add_card(slide, "STATIC CHECKS", "Files • JSON • source wiring • required fields • launch schema", 8.52, 1.85, 3.95, 1.75, BLUE)
    add_card(slide, "HUMAN CHECKS", "Rendered cards • responsive layout • accessibility • browser errors", 8.52, 3.95, 3.95, 1.75, ORANGE)

    section_slide(prs, 5, "Memory, State, and Evaluation", "Preserve useful continuity with provenance while detecting drift and enforcing explicit success signals.", BLUE,
                  """Next we will keep state between steps and check the result
More memory is not always better
Keep the facts that help us continue or recover
Make sure we know where they came from""")

    slide = base_slide(prs, "Three memory horizons", "Module 5 — Memory + evaluation",
                       """Short-term context is for the current task
Long-term memory is for stable, validated information
External state gives us something another person or job can inspect

Ask what should never be retained
Secrets are one example
Also watch for stale context or an assumption being saved as a fact""")
    add_card(slide, "SHORT-TERM", "Current task context\nTool outputs\nWorking hypotheses\nDiscard or compact quickly", 0.78, 1.9, 3.78, 3.45, BLUE)
    add_card(slide, "LONG-TERM", "Stable conventions\nValidated decisions\nReusable preferences\nRetention policy required", 4.78, 1.9, 3.78, 3.45, PURPLE)
    add_card(slide, "EXTERNAL STATE", "Issues + PRs\nArtifacts + logs\nJSONL event history\nSystem of record", 8.78, 1.9, 3.78, 3.45, GREEN)
    add_text(slide, "Persist: facts + provenance. Avoid: secrets, unsupported inference, and unbounded transcripts.", 1.0, 5.95, 11.2, 0.45, 18, ORANGE, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "Event log → snapshot → continuity", "Module 5 — Memory + evaluation",
                       """Run the state_store commands from the lab
Open the event log, then the snapshot
Find the sequence, source, timestamp, and hash

The sequence tells us the order
The source tells us who produced the event
The hash helps detect a change
Keep that provenance when creating a snapshot""")
    state = '''{
  "sequence": 1,
  "event_type": "implementation.accepted",
  "source": "scripts/implement/apply_plan.py",
  "payload": {
    "request_id": "demo-plan-001",
    "status": "ready_for_review",
    "changed_files": ["artifacts/health-report.json"]
  }
}'''
    add_code(slide, state, 0.72, 1.72, 6.4, 4.95, 13, "state event")
    flow(slide, ["APPEND", "ORDER", "HASH", "SNAPSHOT"], 7.48, 2.0, 4.8, 0.95,
         [BLUE, PURPLE, ORANGE, GREEN])
    add_bullets(slide, ["Source identifies producer", "Sequence preserves order", "Hash detects mutation", "Snapshot supports recovery", "Retention prevents stale context"], 7.6, 3.5, 4.55, 2.7, 17)

    comparison_slide(prs, "Evaluation signals and quality gates",
                     "Module 5 — Memory + evaluation",
                     """Some checks have a definite answer
Does the JSON parse? Does the file exist? Do the hashes match?
Other questions need judgment
Is this useful to the user? Is the tradeoff acceptable?

A model score does not replace tests or required review
If a gate fails, report it as a failed outcome
Do not hide it in the handoff""",
                     ("Deterministic checks", ["Schema parses and files exist",
                                               "Hashes match", "Tests pass",
                                               "Risk threshold is respected"]),
                     ("Human judgment", ["Design quality and user value",
                                         "Requirement ambiguity", "Acceptable tradeoffs",
                                         "Novel failure analysis"]))

    section_slide(prs, 6, "Governance, Guardrails, and Operations", "Scale autonomy according to risk, preserve human control, and make every consequential action reversible and auditable.", RED,
                  """Now we will put the controls together
Permissions, approval, evidence, and recovery
Instructions in a prompt are useful, but they are not enforcement
Use GitHub controls for the boundaries that must hold""")

    slide = base_slide(prs, "Risk-based autonomy ladder", "Module 6 — Governance",
                       """Start with reading and summarizing
Then proposing a change, opening a pull request, and deploying
The consequences increase as we move down the slide

Ask learners for one task at each level
Where would they require an approval?
Higher risk needs narrower authority and a clearer recovery path
Some actions should simply be denied""")
    levels = [
        ("1", "READ + SUMMARIZE", "Auto", GREEN),
        ("2", "PROPOSE CHANGE", "Auto + validate", BLUE),
        ("3", "OPEN PULL REQUEST", "Review required", PURPLE),
        ("4", "DEPLOY / MUTATE DATA", "Environment approval", ORANGE),
        ("5", "IRREVERSIBLE ACTION", "Deny or exceptional process", RED),
    ]
    y = 1.65
    for num, action, control, color in levels:
        add_rect(slide, 1.0, y, 11.3, 0.82, fill=PANEL)
        add_text(slide, num, 1.2, y + 0.16, 0.45, 0.35, 18, color, True, align=PP_ALIGN.CENTER)
        add_text(slide, action, 1.9, y + 0.16, 5.2, 0.35, 17, TEXT, True)
        add_text(slide, control, 7.6, y + 0.16, 4.2, 0.35, 16, color, True, align=PP_ALIGN.RIGHT)
        y += 1.0

    slide = base_slide(prs, "Guardrail → approval → audit → recovery", "Module 6 — Governance",
                       """Walk through the request and its permissions
This sample requests a pull request with a protected-branch flag
The permissions are narrow, but the required approval is missing

Run the governance example
The result is approval_required, not execution
Find the audit record and the recovery reference
We need to know who made the decision and which policy was used""")
    flow(slide, ["REQUEST", "POLICY", "APPROVAL", "AUDIT", "ROLLBACK"], 0.75, 1.85, 11.85, 1.05,
         [BLUE, RED, ORANGE, GREEN, PURPLE])
    guard = '''{
  "requested_action": "open_pull_request",
  "permissions": [
    "contents:read",
    "pull-requests:write"
  ],
  "protected_branch": true,
  "approval": {"required": true, "present": false}
}

→ decision: "approval_required"'''
    add_code(slide, guard, 0.95, 3.7, 5.75, 3.0, 12, "governance request")
    add_bullets(slide, ["Narrow required permissions", "Human gate for protected branch", "Run ID + actor + policy version", "Changed paths recorded", "Rollback reference prepared"], 7.15, 3.7, 4.9, 2.6, 17)

    slide = base_slide(prs, "Capstone: prove the whole system", "Workshop capstone",
                       """Bring the pieces together
Submit the artifacts so someone else can inspect and reproduce the work
Not just a message saying the agent succeeded

Use the rubric: bounded, reproducible, observable, reversible
Check the failure path as well as the successful run
Make sure the approval and rollback evidence are included""")
    artifacts = [
        "Agent contract", "Tool + permission policy", "Versioned plan", "Two specialist outputs",
        "Merged plan", "State snapshot", "Evaluation report", "Approval + audit record", "Rollback procedure",
    ]
    for i, item in enumerate(artifacts):
        col, row = i % 3, i // 3
        add_card(slide, item.upper(), "Inspectable evidence", 0.62 + col * 4.2, 1.65 + row * 1.65, 3.85, 1.32,
                 [BLUE, PURPLE, ORANGE, GREEN, BLUE, PURPLE, RED, ORANGE, GREEN][i])
    add_text(slide, "Rubric: bounded • reproducible • observable • reversible", 1.2, 6.65, 10.9, 0.35, 18, TEXT, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "Use the workshop as a GH-600 study map", "Certification preparation",
                       """Use a runnable example for each domain
The contract and handoff for architecture
The tool boundary for permissions
The event log for state, and the fan-in workflow for coordination

Open docs/source-coverage.md
It maps the 52 Microsoft Learn units to the workshop material
Use that map to find the areas that need more practice""")
    mappings = [
        ("Architecture + SDLC", "Agent contract + plan handoff", BLUE),
        ("Tool interaction", "MCP boundary + execution environments", ORANGE),
        ("Memory + state", "Event log + snapshot", PURPLE),
        ("Evaluation", "Quality gates + failure analysis", GREEN),
        ("Multi-agent", "Fan-in workflow + Project Pulse", BLUE),
        ("Guardrails", "Approval + audit + rollback", RED),
    ]
    for i, (domain, evidence, color) in enumerate(mappings):
        col, row = i % 2, i // 2
        x, y = 0.78 + col * 6.18, 1.73 + row * 1.72
        add_rect(slide, x, y, 5.7, 1.32, fill=PANEL)
        add_text(slide, domain, x + 0.25, y + 0.2, 5.1, 0.35, 18, color, True)
        add_text(slide, evidence, x + 0.25, y + 0.7, 5.1, 0.35, 15, MUTED)
    add_text(slide, "docs/source-coverage.md", 4.4, 6.65, 4.5, 0.35, 16, BLUE, True, font=CODE_FONT, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_artwork(slide)
    number = len(prs.slides)
    add_text(slide, "Build agents you can\ninspect, stop, and trust.", 0.85, 2.0, 11.4, 2.0, 46, WHITE, True, TITLE_FONT)
    add_text(slide, "Evidence over autonomy", 0.9, 4.4, 10.9, 0.5, 24, WHITE)
    add_text(slide, "What is one action you would never allow an agent to perform without human approval?",
             0.9, 5.2, 10.9, 0.95, 22, WHITE)
    add_text(slide, "Run the labs  •  inspect the artifacts  •  practice the recovery path",
             0.9, 6.7, 11.5, 0.4, 14, WHITE, font=CODE_FONT)
    add_footer(slide, "Close", number, WHITE)
    add_notes(slide, """Ask for one action they would never allow without human approval
Take a few examples from the room

We want agents we can inspect, interrupt, and recover
Not just agents that work on the happy path
Point learners back to the labs and the study map
Field any final questions""")

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"Generated {output} with {len(prs.slides)} slides")


if __name__ == "__main__":
    make_deck()
