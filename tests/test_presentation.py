from contextlib import redirect_stdout
from io import StringIO
from tempfile import TemporaryDirectory
import unittest
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

from scripts.presentation.generate_workshop_deck import make_deck

REPO = Path(__file__).resolve().parents[1]
DECK = REPO / "slides/github-agentic-ai-developer-workshop.pptx"


class WorkshopPresentationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.presentation = Presentation(DECK)

    def test_deck_is_complete_and_editable(self):
        self.assertTrue(DECK.is_file(), "PowerPoint deck must be generated")
        presentation = self.presentation
        self.assertEqual(len(presentation.slides), 38)

        all_text = []
        slides_with_notes = 0
        for slide in presentation.slides:
            slide_text = "\n".join(
                shape.text for shape in slide.shapes if hasattr(shape, "text")
            )
            all_text.append(slide_text)
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slides_with_notes += 1

        deck_text = "\n".join(all_text)
        for required in (
            "Foundations of Agentic AI",
            "Agent Architecture and SDLC Integration",
            "Tooling, MCP, and Execution Environments",
            "Multi-Agent Systems and Orchestration",
            "Memory, State, and Evaluation",
            "Governance, Guardrails, and Operations",
            "Project Pulse",
            "GH-600",
        ):
            self.assertIn(required, deck_text)
        self.assertNotIn("TODO", deck_text)
        self.assertNotIn("plan=...", deck_text)
        self.assertGreaterEqual(slides_with_notes, len(presentation.slides) - 1)

    def test_github_layouts_and_typography(self):
        presentation = self.presentation
        self.assertAlmostEqual(presentation.slide_width / presentation.slide_height,
                               16 / 9, places=4)
        backgrounds = []
        fonts = set()
        for index, slide in enumerate(presentation.slides):
            with self.subTest(slide=index + 1):
                backgrounds.append(str(slide.background.fill.fore_color.rgb))
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, presentation.slide_width)
                    self.assertLessEqual(shape.top + shape.height, presentation.slide_height)
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            fonts.update(run.font.name or paragraph.font.name
                                         for run in paragraph.runs)
        self.assertEqual(backgrounds.count("1658C5"), 6)
        self.assertEqual(backgrounds.count("FFFFFF"), 30)
        self.assertEqual(backgrounds.count("000000"), 2)
        self.assertEqual(fonts, {"Helvetica Neue", "Roboto Mono"})
        for slide in (presentation.slides[0], presentation.slides[-1]):
            pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
            self.assertEqual(len(pictures), 2, "Title artwork and GitHub mark must be retained")
        for title in ("Assistant versus agent",
                      "Choose the execution environment deliberately",
                      "Evaluation signals and quality gates"):
            slide = next(slide for slide in presentation.slides
                         if any(shape.has_text_frame and shape.text == title
                                for shape in slide.shapes))
            self.assertEqual(str(slide.shapes[0].fill.fore_color.rgb), "F6F8FA")
            self.assertEqual(slide.shapes[0].height, presentation.slide_height)

    def test_notes_use_instructor_cues_and_keep_lab_guidance(self):
        notes = []
        for index, slide in enumerate(self.presentation.slides, 1):
            text = slide.notes_slide.notes_text_frame.text
            with self.subTest(slide=index):
                lines = [line for line in text.splitlines() if line.strip()]
                self.assertGreaterEqual(len(lines), 3, "Use short, separate speaking cues")
                self.assertNotIn("TODO", text)
            notes.append(text)
        all_notes = "\n".join(notes)
        for cue in ("What are you excited to learn today?",
                    "Open .github/workflows/plan-implement.yml",
                    "Check the nonzero exit code",
                    "Select Orchestrator with /agent",
                    "It does not execute the browser JavaScript",
                    "Field any final questions"):
            self.assertIn(cue, all_notes)

    def test_theme_and_bullets_survive_regeneration(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "workshop.pptx"
            with redirect_stdout(StringIO()):
                make_deck(output)
            regenerated = Presentation(output)
            self.assertEqual(len(regenerated.slides), len(self.presentation.slides))
            for original, generated in zip(self.presentation.slides, regenerated.slides):
                self.assertEqual(
                    [shape.text for shape in original.shapes if shape.has_text_frame],
                    [shape.text for shape in generated.shapes if shape.has_text_frame],
                )
                self.assertEqual(original.notes_slide.notes_text_frame.text,
                                 generated.notes_slide.notes_text_frame.text)
            with ZipFile(output) as archive:
                theme = etree.fromstring(archive.read("ppt/theme/theme1.xml"))
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                self.assertEqual(theme.get("name"), "GitHub Developer Training")
                self.assertEqual(
                    theme.xpath(".//a:fontScheme/*/a:latin/@typeface", namespaces=ns),
                    ["Helvetica Neue", "Helvetica Neue"],
                )
                self.assertEqual(
                    theme.xpath(".//a:clrScheme/a:accent1/a:srgbClr/@val", namespaces=ns),
                    ["0366D6"],
                )
                self.assertFalse(theme.xpath(".//a:effectStyle/a:effectLst/*", namespaces=ns))
                bullets = sum(
                    len(etree.fromstring(archive.read(name)).xpath(".//a:buChar", namespaces=ns))
                    for name in archive.namelist()
                    if name.startswith("ppt/slides/slide") and name.endswith(".xml")
                )
                self.assertGreater(bullets, 20)


if __name__ == "__main__":
    unittest.main()
